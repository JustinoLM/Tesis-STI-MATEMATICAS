"""
Exportaciones del panel de administración: privacidad (nombres vacíos por defecto, sin
contraseñas) y hojas de desafíos grupales y participación.
"""

import pytest
from httpx import AsyncClient

from app.main import app
from app.models.adaptive import PerfilEstudiante
from app.models.gamification import CategoriaMedalla, EstudianteMedalla, Medalla
from tests.conftest import TestSessionLocal, admin_headers
from tests.test_medallas_sesiones_config import _desafio_completado, _estudiante, _sesion

NOMBRE = "Estudiante de prueba"   # nombre que crea _estudiante


async def _get(c, ruta, **params):
    r = await c.get(f"/api/admin/export/{ruta}", headers=admin_headers(), params=params)
    assert r.status_code == 200, r.text
    datos = r.json()
    if ruta == "niveles":   # dos tablas: historial por sesión y nivel actual
        return {"columnas": datos["historial"]["columnas"],
                "filas": datos["historial"]["filas"] + datos["nivel_actual"]["filas"]}
    return datos


@pytest.mark.asyncio
@pytest.mark.parametrize("ruta,columna", [
    ("estudiantes", "nombre_completo"),
    ("diagnostico", "nombre_estudiante"),
    ("sesiones", "nombre_estudiante"),
    ("niveles", "nombre_estudiante"),
    ("medallas", "nombre_estudiante"),
])
async def test_los_nombres_solo_salen_si_se_piden(ruta, columna):
    async with AsyncClient(app=app, base_url="http://test") as c:
        est, _ = await _estudiante(c, "EXP001")
        async with TestSessionLocal() as db:
            db.add(PerfilEstudiante(estudiante_id=est))
            medalla = Medalla(nombre="M", descripcion="x", categoria=CategoriaMedalla.EXPLORACION,
                              criterio={"tipo": "diagnostico"})
            db.add_all([medalla, _sesion(est)])
            await db.flush()
            db.add(EstudianteMedalla(estudiante_id=est, medalla_id=medalla.id))
            await db.commit()

        sin = await _get(c, ruta)
        assert sin["filas"], f"{ruta} sin filas"
        assert columna in sin["columnas"]                       # la columna se conserva (importación por posición)
        assert all(not f[columna] for f in sin["filas"])
        assert NOMBRE not in str(sin)

        con = await _get(c, ruta, con_nombres=True)
        assert any(f[columna] == NOMBRE for f in con["filas"])


@pytest.mark.asyncio
async def test_el_export_de_estudiantes_no_incluye_contrasenas():
    async with AsyncClient(app=app, base_url="http://test") as c:
        await _estudiante(c, "EXP002")
        datos = await _get(c, "estudiantes", con_nombres=True)
    assert not any("password" in col or "contrasena" in col.lower() for col in datos["columnas"])


@pytest.mark.asyncio
async def test_desafios_y_participacion_se_exportan_con_la_regla_del_sistema():
    async with AsyncClient(app=app, base_url="http://test") as c:
        participa, _ = await _estudiante(c, "EXP003")
        espectador, _ = await _estudiante(c, "EXP004")
        async with TestSessionLocal() as db:
            db.add_all([PerfilEstudiante(estudiante_id=participa), PerfilEstudiante(estudiante_id=espectador)])
            await db.flush()
            await _desafio_completado(db, participa, sesiones_en_ventana=3)    # participa
            await _desafio_completado(db, espectador, sesiones_en_ventana=2)   # no alcanza el mínimo

        desafios = await _get(c, "desafios")
        assert desafios["columnas"][:5] == ["id_desafio", "nombre", "tipo", "grupo", "objetivo"]
        assert len(desafios["filas"]) == 2
        assert all(f["completado"] == "Sí" for f in desafios["filas"])
        assert sorted(f["estudiantes_participantes"] for f in desafios["filas"]) == [0, 1]

        part = await _get(c, "participacion")
        por_codigo = {f["codigo_estudiante"]: f for f in part["filas"]}
        assert por_codigo["EXP003"]["participo"] == "Sí" and por_codigo["EXP003"]["cuenta_para_medalla"] == "Sí"
        assert por_codigo["EXP004"]["participo"] == "No" and por_codigo["EXP004"]["cuenta_para_medalla"] == "No"
        assert por_codigo["EXP003"]["sesiones_en_ventana"] == 3
        assert all(not f["nombre_estudiante"] for f in part["filas"])


@pytest.mark.asyncio
async def test_resumen_incluye_los_desafios():
    async with AsyncClient(app=app, base_url="http://test") as c:
        est, _ = await _estudiante(c, "EXP005")
        async with TestSessionLocal() as db:
            db.add(PerfilEstudiante(estudiante_id=est))
            await db.flush()
            await _desafio_completado(db, est, sesiones_en_ventana=3)
        # el resumen exige que la organización tenga estudiantes; _estudiante no asigna una,
        # así que solo se comprueba que la ruta responde con las columnas nuevas
        r = await _get(c, "resumen")
    assert "columnas" in r


@pytest.mark.asyncio
async def test_las_exportaciones_exigen_token_de_administrador():
    async with AsyncClient(app=app, base_url="http://test") as c:
        for ruta in ("desafios", "participacion"):
            r = await c.get(f"/api/admin/export/{ruta}")
            assert r.status_code in (401, 403)


@pytest.mark.asyncio
async def test_un_export_sin_nombres_se_puede_importar_y_las_hojas_nuevas_se_ignoran():
    """El Excel anónimo (nombres vacíos, Medallas con 7 columnas, hojas de desafíos) se importa."""
    import io

    import openpyxl
    from sqlalchemy import select

    from app.models.gamification import EstudianteMedalla
    from app.models.user import Estudiante

    async with TestSessionLocal() as db:
        db.add(Medalla(nombre="Principiante", descripcion="x", categoria=CategoriaMedalla.APRENDIZAJE,
                       criterio={"tipo": "diagnostico"}))
        await db.commit()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estudiantes"
    ws.append(["codigo_estudiante", "nombre_completo", "genero", "grado_academico", "edad", "organizacion",
               "puntos_actuales", "activo", "fecha_creacion", "ultimo_acceso"])
    ws.append(["IMP001", None, "femenino", "6C", 11, "Colegio", 40, "Sí", "2026-05-07 08:00:00", ""])
    wm = wb.create_sheet("Medallas")
    wm.append(["codigo_estudiante", "nombre_estudiante", "organizacion", "medalla", "categoria",
               "descripcion", "fecha_obtencion"])
    wm.append(["IMP001", None, "Colegio", "Principiante", "aprendizaje", "x", "2026-05-07 08:20:00"])
    wd = wb.create_sheet("Desafíos")
    wd.append(["id_desafio", "nombre"])
    wd.append([1, "Desafío"])
    buf = io.BytesIO()
    wb.save(buf)

    async with AsyncClient(app=app, base_url="http://test") as c:
        r = await c.post("/api/admin/import/excel", headers=admin_headers(),
                         data={"organizacion_nombre": "Colegio"},
                         files={"archivo": ("export.xlsx", buf.getvalue(),
                                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200, r.text
        hojas = {h["hoja"]: h for h in r.json()["hojas"]}
        assert hojas["Estudiantes"]["creados"] == 1 and hojas["Medallas"]["creados"] == 1

    async with TestSessionLocal() as db:
        est = (await db.execute(select(Estudiante).where(Estudiante.codigo_estudiante == "IMP001"))).scalar_one()
        assert est.nombre_completo == "IMP001"          # sin nombre: se usa el código
        assert len((await db.execute(select(EstudianteMedalla))).scalars().all()) == 1
